"""Documenten: teksten maken en bewerken in de app, op elk apparaat met een browser (iPad, iPhone, Mac, Windows).

De tekst staat in de database, dus wat je op de iPad schrijft, staat meteen ook op je laptop. Elk document heeft een
versienummer: is het intussen op een ander apparaat gewijzigd, dan overschrijft de app dat niet zomaar.
Hier staat ook het omzetten van en naar Word, PDF en gewone tekst.
"""

import base64
import html
import io
import os
import re
from html.parser import HTMLParser

import opslag

# Lettertypen voor PDF met alle tekens (ook €). Op Streamlit Cloud via packages.txt (fonts-dejavu-core).
LETTERS = "/usr/share/fonts/truetype/dejavu/DejaVuSans{}.ttf"


class Fout(Exception):
    """Iets lukte niet; de tekst is geschikt om aan de gebruiker te tonen."""


class Conflict(Exception):
    """Het document is intussen ergens anders gewijzigd. `document` is de nieuwste versie."""

    def __init__(self, document: dict):
        super().__init__("Dit document is intussen op een ander apparaat gewijzigd.")
        self.document = document


def _ok(antwoord: dict) -> dict:
    if antwoord.get("fout"):
        raise Fout(antwoord["fout"])
    if antwoord.get("conflict"):
        raise Conflict(antwoord["document"])
    return antwoord


# --- Opslag ---------------------------------------------------------------------------------------------------------

def lijst(zoek: str = "", prullenbak: bool = False) -> list[dict]:
    return opslag._rpc_ik("documenten_lijst", p_zoek=zoek, p_prullenbak=prullenbak)["documenten"]


def lezen(document_id: int) -> dict:
    return _ok(opslag._rpc_ik("document_lezen", p_id=document_id))["document"]


def opslaan(document_id: int | None, titel: str | None = None, inhoud: dict | None = None, html_tekst: str | None = None,
            versie: int | None = None, apparaat: str = "") -> dict:
    """Bewaart een document (nieuw als document_id leeg is). Geeft id en nieuwe versie; Conflict bij een botsing.

    Zonder `inhoud` maar met `html_tekst` bouwt de editor de tekst bij het openen op uit de html.
    """
    tekst = None if html_tekst is None else naar_tekst(html_tekst)
    return _ok(opslag._rpc_ik("document_opslaan", p_id=document_id, p_titel=titel, p_inhoud=inhoud,
                              p_html=html_tekst, p_tekst=tekst, p_versie=versie, p_apparaat=apparaat))["document"]


def maken(titel: str, html_tekst: str = "", apparaat: str = "") -> dict:
    return opslaan(None, titel, None, html_tekst, None, apparaat)


def aanvullen(document_id: int, html_tekst: str, apparaat: str = "") -> dict:
    """Zet tekst onderaan een document (voor de assistent). De editor bouwt de tekst daarna op uit de html."""
    huidig = lezen(document_id)
    return opslaan(document_id, None, None, (huidig.get("html") or "") + html_tekst, huidig["versie"], apparaat)


def weg(document_id: int, naar_prullenbak: bool = True) -> None:
    _ok(opslag._rpc_ik("document_weg", p_id=document_id, p_weg=naar_prullenbak))


def wissen(document_id: int) -> None:
    _ok(opslag._rpc_ik("document_wissen", p_id=document_id))


def versies(document_id: int) -> list[dict]:
    return opslag._rpc_ik("document_versies", p_id=document_id)["versies"]


def terugzetten(versie_id: int, apparaat: str = "") -> dict:
    return _ok(opslag._rpc_ik("document_versie_terugzetten", p_versie_id=versie_id, p_apparaat=apparaat))["document"]


# --- Tekst en html ------------------------------------------------------------------------------------------------

_BLOK = re.compile(r"</(p|h[1-6]|li|blockquote|pre|div)>|<br\s*/?>", re.I)


def naar_tekst(html_tekst: str) -> str:
    """Platte tekst uit html: blokken op een eigen regel, zonder opmaak."""
    tekst = _BLOK.sub("\n", html_tekst or "")
    tekst = re.sub(r"<li[^>]*>", "• ", tekst, flags=re.I)
    tekst = html.unescape(re.sub(r"<[^>]+>", "", tekst)).replace("\xa0", " ")
    return re.sub(r"\n{3,}", "\n\n", tekst).strip()


def _inline_md(regel: str) -> str:
    regel = html.escape(regel, quote=False)
    regel = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", regel)
    regel = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"<em>\1</em>", regel)
    regel = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', regel)
    return regel


def markdown_naar_html(tekst: str) -> str:
    """Eenvoudige markdown (kopjes, lijstjes, vet, schuin, links) naar html voor de editor."""
    uit, lijst_soort = [], None

    def sluit_lijst():
        nonlocal lijst_soort
        if lijst_soort:
            uit.append(f"</{lijst_soort}>")
            lijst_soort = None

    for regel in (tekst or "").replace("\r", "").split("\n"):
        kaal = regel.strip()
        if m := re.match(r"^(#{1,3})\s+(.*)", kaal):
            sluit_lijst()
            n = len(m.group(1))
            uit.append(f"<h{n}>{_inline_md(m.group(2))}</h{n}>")
        elif m := re.match(r"^[-*•]\s+(.*)", kaal):
            if lijst_soort != "ul":
                sluit_lijst()
                uit.append("<ul>")
                lijst_soort = "ul"
            uit.append(f"<li>{_inline_md(m.group(1))}</li>")
        elif m := re.match(r"^\d+[.)]\s+(.*)", kaal):
            if lijst_soort != "ol":
                sluit_lijst()
                uit.append("<ol>")
                lijst_soort = "ol"
            uit.append(f"<li>{_inline_md(m.group(1))}</li>")
        elif m := re.match(r"^>\s?(.*)", kaal):
            sluit_lijst()
            uit.append(f"<blockquote>{_inline_md(m.group(1))}</blockquote>")
        elif kaal:
            sluit_lijst()
            uit.append(f"<p>{_inline_md(kaal)}</p>")
        else:
            sluit_lijst()
    sluit_lijst()
    return "".join(uit)


def tekst_naar_html(tekst: str) -> str:
    return "".join(f"<p>{html.escape(r.strip(), quote=False)}</p>" for r in re.split(r"\n\s*\n|\n", tekst or "")
                   if r.strip())


# --- Importeren ---------------------------------------------------------------------------------------------------

def importeren(naam: str, inhoud: bytes) -> tuple[str, str]:
    """Een bestand (Word, PDF, tekst of markdown) als (titel, html)."""
    titel = os.path.splitext(os.path.basename(naam))[0] or "Geïmporteerd document"
    soort = os.path.splitext(naam)[1].lower()
    if soort == ".docx":
        return titel, _docx_naar_html(inhoud)
    if soort == ".pdf":
        from pypdf import PdfReader
        lezer = PdfReader(io.BytesIO(inhoud))
        return titel, "".join(tekst_naar_html(p.extract_text() or "") for p in lezer.pages)
    tekst = inhoud.decode("utf-8", errors="replace")
    if soort in (".html", ".htm"):
        return titel, re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", "", tekst)
    return titel, markdown_naar_html(tekst) if soort == ".md" else tekst_naar_html(tekst)


def _docx_naar_html(inhoud: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(inhoud))
    uit, lijst_soort = [], None
    for alinea in doc.paragraphs:
        stijl = (alinea.style.name or "").lower() if alinea.style is not None else ""
        delen = []
        for run in alinea.runs:
            stuk = html.escape(run.text, quote=False)
            if not stuk:
                continue
            if run.bold:
                stuk = f"<strong>{stuk}</strong>"
            if run.italic:
                stuk = f"<em>{stuk}</em>"
            if run.underline:
                stuk = f"<u>{stuk}</u>"
            delen.append(stuk)
        tekst = "".join(delen)
        soort = "ol" if "number" in stijl else "ul" if ("bullet" in stijl or "list" in stijl) else None
        if soort != lijst_soort:
            if lijst_soort:
                uit.append(f"</{lijst_soort}>")
            if soort:
                uit.append(f"<{soort}>")
            lijst_soort = soort
        if soort:
            uit.append(f"<li>{tekst}</li>")
        elif m := re.match(r"(heading|kop)\s*(\d)", stijl):
            n = min(int(m.group(2)), 3)
            uit.append(f"<h{n}>{tekst}</h{n}>")
        elif stijl == "title":
            uit.append(f"<h1>{tekst}</h1>")
        else:
            uit.append(f"<p>{tekst}</p>" if tekst else "<p><br></p>")
    if lijst_soort:
        uit.append(f"</{lijst_soort}>")
    for tabel in doc.tables:
        for rij in tabel.rows:
            uit.append("<p>" + " | ".join(html.escape(c.text, quote=False) for c in rij.cells) + "</p>")
    return "".join(uit)


# --- Exporteren ---------------------------------------------------------------------------------------------------

class _Lezer(HTMLParser):
    """Leest de html van de editor als een lijst blokken met opgemaakte stukjes tekst."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blokken: list[dict] = []
        self.opmaak: list[str] = []
        self.lijsten: list[list] = []  # [soort, teller]
        self.blok: dict | None = None
        self.link: str | None = None

    def _nieuw(self, soort: str, **extra) -> None:
        self.blok = {"soort": soort, "stukken": [], **extra}
        self.blokken.append(self.blok)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre"):
            self._nieuw(tag, uitlijning=_uitlijning(attrs))
        elif tag in ("ul", "ol"):
            self.lijsten.append([tag, 0])
        elif tag == "li":
            soort = self.lijsten[-1][0] if self.lijsten else "ul"
            if self.lijsten:
                self.lijsten[-1][1] += 1
            vink = attrs.get("data-list")
            self._nieuw("li", lijst=soort, nummer=self.lijsten[-1][1] if self.lijsten else 1,
                        diepte=max(len(self.lijsten) - 1, 0), vink=vink if vink in ("checked", "unchecked") else None)
        elif tag in ("strong", "b", "em", "i", "u", "s", "strike", "code", "sub", "sup"):
            self.opmaak.append({"b": "strong", "i": "em", "strike": "s"}.get(tag, tag))
        elif tag == "a":
            self.link = attrs.get("href")
        elif tag == "br":
            if self.blok is None:
                self._nieuw("p")
            self.blok["stukken"].append({"tekst": "\n", "opmaak": set()})
        elif tag == "img" and (bron := attrs.get("src", "")).startswith("data:image"):
            self._nieuw("img", bron=bron)
            self.blok = None

    def handle_endtag(self, tag):
        if tag in ("strong", "b", "em", "i", "u", "s", "strike", "code", "sub", "sup"):
            naam = {"b": "strong", "i": "em", "strike": "s"}.get(tag, tag)
            if naam in self.opmaak:
                self.opmaak.reverse()
                self.opmaak.remove(naam)
                self.opmaak.reverse()
        elif tag == "a":
            self.link = None
        elif tag in ("ul", "ol") and self.lijsten:
            self.lijsten.pop()
        elif tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre", "li"):
            self.blok = None

    def handle_data(self, data):
        if not data.strip("\n"):
            return
        if self.blok is None:
            if not data.strip():
                return
            self._nieuw("p")
        self.blok["stukken"].append({"tekst": data.replace("\xa0", " "), "opmaak": set(self.opmaak),
                                     "link": self.link})


def _uitlijning(attrs: dict) -> str | None:
    klasse = attrs.get("class") or ""
    stijl = attrs.get("style") or ""
    for waarde in ("center", "right", "justify"):
        if f"ql-align-{waarde}" in klasse or f"text-align: {waarde}" in stijl or f"text-align:{waarde}" in stijl:
            return waarde
    return None


def _blokken(html_tekst: str) -> list[dict]:
    lezer = _Lezer()
    lezer.feed(html_tekst or "")
    return [b for b in lezer.blokken if b["soort"] == "img" or any(s["tekst"].strip() for s in b["stukken"])
            or b["soort"] == "p"]


def _afbeelding(bron: str) -> io.BytesIO | None:
    try:
        return io.BytesIO(base64.b64decode(bron.split(",", 1)[1]))
    except (IndexError, ValueError):
        return None


def naar_word(titel: str, html_tekst: str) -> bytes:
    """Het document als Word-bestand (.docx): opent in Word, Pages en Google Docs."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    doc.core_properties.title = titel
    stijl = doc.styles["Normal"]
    stijl.font.name = "Calibri"
    stijl.font.size = Pt(11)
    uitlijnen = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                 "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    for blok in _blokken(html_tekst):
        soort = blok["soort"]
        if soort == "img":
            if beeld := _afbeelding(blok["bron"]):
                try:
                    doc.add_picture(beeld, width=Cm(15))
                except Exception:  # onbekend beeldformaat: overslaan
                    pass
            continue
        if soort in ("h1", "h2", "h3", "h4", "h5", "h6"):
            alinea = doc.add_heading(level=min(int(soort[1]), 3))
        elif soort == "li":
            if blok.get("vink"):
                alinea = doc.add_paragraph()
                alinea.add_run("☑ " if blok["vink"] == "checked" else "☐ ")
            else:
                naam = "List Number" if blok["lijst"] == "ol" else "List Bullet"
                naam += f" {blok['diepte'] + 1}" if blok["diepte"] else ""
                try:
                    alinea = doc.add_paragraph(style=naam)
                except KeyError:
                    alinea = doc.add_paragraph(style="List Bullet")
        elif soort == "blockquote":
            alinea = doc.add_paragraph(style="Quote")
        else:
            alinea = doc.add_paragraph()
        if blok.get("uitlijning") in uitlijnen:
            alinea.alignment = uitlijnen[blok["uitlijning"]]
        for stuk in blok["stukken"]:
            run = alinea.add_run(stuk["tekst"])
            run.bold = "strong" in stuk["opmaak"] or None
            run.italic = "em" in stuk["opmaak"] or None
            run.underline = "u" in stuk["opmaak"] or None
            run.font.strike = "s" in stuk["opmaak"] or None
            if "code" in stuk["opmaak"] or soort == "pre":
                run.font.name = "Courier New"
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def naar_pdf(titel: str, html_tekst: str) -> bytes:
    """Het document als PDF, om te printen of te versturen."""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=20)
    pdf.set_title(titel)
    if os.path.exists(LETTERS.format("")):
        for stijl, achtervoegsel in (("", ""), ("B", "-Bold"), ("I", "-Oblique"), ("BI", "-BoldOblique")):
            pad = LETTERS.format(achtervoegsel)
            pdf.add_font("Tekst", stijl, pad if os.path.exists(pad) else LETTERS.format(""))
        letter, veilig = "Tekst", (lambda t: t)
    else:
        # Zonder eigen lettertype kent PDF alleen westerse tekens.
        letter = "Helvetica"
        veilig = (lambda t: t.replace("€", "EUR").replace("’", "'").replace("‘", "'").replace("“", '"')
                  .replace("”", '"').replace("–", "-").replace("—", "-").replace("•", "-").replace("☑", "[x]")
                  .replace("☐", "[ ]").encode("latin-1", "replace").decode("latin-1"))
    pdf.add_page()
    groottes = {"h1": 20, "h2": 16, "h3": 13, "h4": 12, "h5": 11, "h6": 11}
    uitlijnen = {"center": "C", "right": "R", "justify": "J"}
    for blok in _blokken(html_tekst):
        soort = blok["soort"]
        if soort == "img":
            if beeld := _afbeelding(blok["bron"]):
                try:
                    pdf.image(beeld, w=pdf.epw * 0.9)
                    pdf.ln(3)
                except Exception:
                    pass
            continue
        grootte = groottes.get(soort, 11)
        inspring = 0
        voor = ""
        if soort == "li":
            inspring = 6 + 6 * blok["diepte"]
            voor = ("☑ " if blok["vink"] == "checked" else "☐ ") if blok.get("vink") else (
                f"{blok['nummer']}. " if blok["lijst"] == "ol" else "• ")
        elif soort == "blockquote":
            inspring = 8
        hoogte = grootte * 0.5
        pdf.set_left_margin(20 + inspring)
        pdf.set_x(20 + inspring)

        def stijl(stuk: dict) -> str:
            return (("B" if "strong" in stuk["opmaak"] or soort.startswith("h") else "")
                    + ("I" if "em" in stuk["opmaak"] or soort == "blockquote" else "")
                    + ("U" if "u" in stuk["opmaak"] or stuk.get("link") else ""))

        stijlen = {stijl(s) for s in blok["stukken"]}
        if blok.get("uitlijning") in uitlijnen and len(stijlen) == 1:
            # Gecentreerd of uitgevuld kan alleen met één soort opmaak in de hele alinea.
            pdf.set_font(letter, stijlen.pop(), grootte)
            pdf.multi_cell(0, hoogte, veilig(voor + "".join(s["tekst"] for s in blok["stukken"])),
                           align=uitlijnen[blok["uitlijning"]], new_x="LMARGIN", new_y="NEXT")
        else:
            if voor:
                pdf.set_font(letter, "", grootte)
                pdf.write(hoogte, veilig(voor))
            for stuk in blok["stukken"]:
                pdf.set_font(letter, stijl(stuk), grootte)
                pdf.write(hoogte, veilig(stuk["tekst"]), link=stuk.get("link") or "")
            pdf.ln(hoogte)
        pdf.set_left_margin(20)
        pdf.ln(4 if soort.startswith("h") else 1.5 if soort == "li" else 2.5)
    return bytes(pdf.output())


def bestandsnaam(titel: str, extensie: str) -> str:
    naam = re.sub(r"[^\w\- ]+", "", titel or "").strip() or "document"
    return f"{naam[:60]}.{extensie}"


def stand(document_id: int) -> dict:
    """Versienummer en waar het het laatst is gewijzigd, zonder de hele tekst op te halen."""
    return _ok(opslag._rpc_ik("document_stand", p_id=document_id))
